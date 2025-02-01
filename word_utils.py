from docx import Document

def load_word_template(vorlage_pfad):
    """Lädt eine Word-Vorlage und gibt ein Dokument-Objekt zurück."""
    return Document(vorlage_pfad)

def replace_placeholders(document, daten):
    """Ersetzt Platzhalter im Word-Dokument mit den entsprechenden Werten aus `daten`."""
    for paragraph in document.paragraphs:
        for key, value in daten.items():
            placeholder = f"[{key}]"
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))

def save_word_document(document, speicher_pfad):
    """Speichert das bearbeitete Word-Dokument."""
    document.save(speicher_pfad)
